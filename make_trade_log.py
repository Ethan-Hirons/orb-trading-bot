"""Regenerate trade_log.xlsx from state/trade_history.csv. Read-only vs the bot.

Usage:  python make_trade_log.py     (needs: pip install openpyxl)
Safe to rerun any time; overwrites trade_log.xlsx in the repo root.
"""
import csv
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent
OUT = REPO / "trade_log.xlsx"

# Starting equity used for the return-on-capital rows. Override per account:
#   START_EQUITY=25000 python make_trade_log.py
START_EQUITY = float(os.environ.get("START_EQUITY", "10000"))

# Exit types classified from logs (see bookkeeping.py); new trades left blank —
# the daily recap fills them in, or run bookkeeping.py for the classification.
EXIT_TYPES = {
    ("2026-07-09", "TQQQ"): "eod-flatten",
    ("2026-07-10", "TRAX"): "stop",
    ("2026-07-13", "NVVE"): "eod-flatten",
    ("2026-07-16", "PYPG"): "take-profit",
    ("2026-07-17", "ATAI"): "eod-flatten",
    ("2026-07-17", "INTC"): "overnight carry (v1.4 bug)",
    ("2026-07-20", "SOXS"): "time-stop",
    ("2026-07-20", "NOK"): "time-stop",
    ("2026-07-20", "ONDS"): "time-stop",
    ("2026-07-21", "IREG"): "time-stop",
    ("2026-07-21", "IREX"): "time-stop",
    ("2026-07-21", "IRE"): "time-stop",
    ("2026-07-22", "NBIL"): "stop",
    ("2026-07-22", "CLBK"): "time-stop",
    ("2026-07-23", "SKYQ"): "eod-flatten",
    ("2026-07-23", "SMCX"): "stop",
    ("2026-07-23", "SMCL"): "stop",
    ("2026-07-23", "SNXX"): "stop",
    ("2026-07-24", "INTC"): "news-exit",
    ("2026-07-24", "TQQQ"): "time-stop",
    ("2026-07-27", "NOK"): "time-stop",
    ("2026-07-27", "SAFT"): "time-stop",
    ("2026-07-28", "CAPR"): "take-profit",
    ("2026-07-28", "DRAM"): "stop",
    ("2026-07-28", "INTC"): "stop",
    ("2026-07-28", "NOK"): "stop",
    ("2026-07-28", "FBRX"): "time-stop",
    ("2026-07-28", "BITO"): "time-stop",
    ("2026-07-29", "REPL"): "take-profit",
    ("2026-07-29", "AAL"): "time-stop",
    ("2026-07-30", "SNXX"): "take-profit",
    ("2026-07-30", "DRAM"): "time-stop",
    ("2026-07-30", "INTC"): "time-stop",
    ("2026-07-30", "NVDA"): "time-stop",
    ("2026-07-30", "PSN"): "time-stop",
    ("2026-07-31", "NBIZ"): "stop",
    ("2026-07-31", "SNDQ"): "stop",
    ("2026-07-31", "AMZN"): "time-stop",
    ("2026-08-03", "SNXX"): "breaker",
    ("2026-08-03", "INTC"): "breaker",
    ("2026-08-03", "SOXL"): "breaker",
    ("2026-08-04", "SOXL"): "stop",
    ("2026-08-04", "SNAP"): "stop",
    ("2026-08-04", "CWVX"): "stop",
    ("2026-08-04", "INTC"): "eod-flatten",
    ("2026-08-05", "SNXX"): "stop",
    ("2026-08-05", "INTC"): "stop",
    ("2026-08-05", "PLTZ"): "stop",
    ("2026-08-05", "SPCX"): "stop",
    ("2026-08-05", "SPCX"): "stop",
    ("2026-08-05", "AHCO"): "stop",
    ("2026-08-06", "INSM"): "stop",
    ("2026-08-06", "WPP"): "stop",
    ("2026-08-11", "ACHR"): "stop",
    ("2026-08-12", "SNXX"): "stop",
    ("2026-08-13", "NEBX"): "stop",
    ("2026-08-13", "ONDS"): "stop",
    ("2026-08-13", "SMCI"): "trail",
    ("2026-08-13", "CLBT"): "stop",
    ("2026-08-13", "INTC"): "trail",
    ("2026-08-14", "TSLL"): "stop",
    ("2026-08-14", "CLBT"): "stop",
    ("2026-08-19", "SOXS"): "trail",
    ("2026-08-19", "AMLX"): "trail",
    ("2026-08-20", "BITO"): "eod-flatten",
    ("2026-08-21", "MARA"): "stop",
    ("2026-08-21", "CONL"): "stop",
    ("2026-08-21", "TSLL"): "trail",
    ("2026-08-21", "MRNA"): "trail",
    ("2026-08-21", "ETHA"): "stop",
    ("2026-08-21", "BITO"): "stop",
    ("2026-08-21", "IBIT"): "eod-flatten",
    ("2026-08-24", "SOXS"): "stop",
    ("2026-08-24", "BMNR"): "trail",
    ("2026-08-24", "ETHA"): "stop",
    ("2026-08-24", "IBIT"): "eod-flatten",
    ("2026-08-24", "BITO"): "eod-flatten",
    ("2026-08-25", "NVDA"): "eod-flatten",
    ("2026-08-26", "SNXX"): "stop",
    ("2026-08-26", "DKS"): "trail",
    ("2026-08-26", "PATH"): "stop",
    ("2026-09-08", "DYN"): "stop",
    ("2026-09-08", "TSLL"): "trail",
    ("2026-09-09", "IRD"): "trail",
    ("2026-09-09", "INTC"): "trail",
}

rows = list(csv.DictReader(open(REPO / "state" / "trade_history.csv", newline="")))

ARIAL = "Arial"
F = lambda **kw: Font(name=ARIAL, **{"size": 10, **kw})
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FONT = F(color="0000FF")  # blue = hardcoded input
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
ws = wb.active
ws.title = "Trade Log"

headers = ["Date", "Symbol", "Side", "Qty", "Entry $", "Exit $", "P&L $",
           "P&L %", "Win/Loss", "Cumulative P&L $", "Exit type"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = F(bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center")
    cell.border = BORDER

# Pre-filled formula rows. Must never be smaller than the number of trades in
# state/trade_history.csv, or trades fall silently outside every formula range
# (this happened: the sheet capped at 60 while the CSV had 65, so the Summary
# read +4.67 instead of +8.86). Grows in blocks of 20 with headroom.
N_ROWS = max(60, ((len(rows) + 20) // 20) * 20)
LAST = N_ROWS + 1  # last data row on the Trade Log sheet
assert LAST - 1 >= len(rows), f"row cap {N_ROWS} < {len(rows)} trades"
for i in range(N_ROWS):
    r = i + 2
    if i < len(rows):
        row = rows[i]
        ws.cell(row=r, column=1, value=row["date"]).font = INPUT_FONT
        ws.cell(row=r, column=2, value=row["symbol"]).font = INPUT_FONT
        ws.cell(row=r, column=3, value=row["side"]).font = INPUT_FONT
        ws.cell(row=r, column=4, value=float(row["qty"])).font = INPUT_FONT
        ws.cell(row=r, column=5, value=float(row["entry_price"])).font = INPUT_FONT
        ws.cell(row=r, column=6, value=float(row["exit_price"])).font = INPUT_FONT
        ws.cell(row=r, column=11,
                value=EXIT_TYPES.get((row["date"], row["symbol"]), "")).font = INPUT_FONT
    ws.cell(row=r, column=7,
            value=f'=IF($B{r}="","",IF($C{r}="long",($F{r}-$E{r})*$D{r},($E{r}-$F{r})*$D{r}))').font = F()
    ws.cell(row=r, column=8,
            value=f'=IF($B{r}="","",IF($C{r}="long",$F{r}/$E{r}-1,1-$F{r}/$E{r}))').font = F()
    ws.cell(row=r, column=9,
            value=f'=IF($B{r}="","",IF($G{r}>=0,"WIN","LOSS"))').font = F(bold=True)
    ws.cell(row=r, column=10,
            value=f'=IF($B{r}="","",SUM($G$2:$G{r}))').font = F()
    for c in range(1, 12):
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=8).number_format = "0.00%"
    for c in (5, 6, 7, 10):
        ws.cell(row=r, column=c).number_format = '#,##0.00;(#,##0.00);"-"'

# --- Grand-total row beneath the data (immediately below the last data row) ---
TOT = LAST + 1
tl = ws.cell(row=TOT, column=6, value="TOTAL")
tl.font = F(bold=True)
tl.alignment = Alignment(horizontal="right")
tg = ws.cell(row=TOT, column=7, value=f"=SUM($G$2:$G${N_ROWS + 1})")
tg.font = F(bold=True)
tg.number_format = '#,##0.00;(#,##0.00);"-"'
# net return on starting equity, shown beside the total
tr = ws.cell(row=TOT, column=8, value=f"=$G${TOT}/{START_EQUITY:g}")
tr.font = F(bold=True)
tr.number_format = "0.00%"
tw = ws.cell(row=TOT, column=9, value="net")
tw.font = F(bold=True)
top = Side(style="double", color="1F3864")
for c in range(1, 12):
    ws.cell(row=TOT, column=c).border = Border(top=top)

widths = [11, 9, 7, 6, 9, 9, 10, 9, 10, 15, 24]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"

s = wb.create_sheet("Summary")
T = "'Trade Log'"
items = [
    ("PAPER TRIAL SUMMARY", None, None),
    ("Completed trades", f"=COUNTIF({T}!$B$2:$B${LAST},\"<>\")", "0"),
    ("Wins", f"=COUNTIF({T}!$I$2:$I${LAST},\"WIN\")", "0"),
    ("Losses", f"=COUNTIF({T}!$I$2:$I${LAST},\"LOSS\")", "0"),
    ("Win rate", f"=IF($B$2=0,\"\",$B$3/$B$2)", "0.0%"),
    ("Net P&L $", f"=SUM({T}!$G$2:$G${LAST})", '#,##0.00;(#,##0.00);"-"'),
    ("Avg win $", f"=IFERROR(AVERAGEIF({T}!$I$2:$I${LAST},\"WIN\",{T}!$G$2:$G${LAST}),\"\")", '#,##0.00'),
    ("Avg loss $", f"=IFERROR(AVERAGEIF({T}!$I$2:$I${LAST},\"LOSS\",{T}!$G$2:$G${LAST}),\"\")", '#,##0.00;(#,##0.00)'),
    ("Expectancy $/trade", f"=IF($B$2=0,\"\",$B$6/$B$2)", '#,##0.00;(#,##0.00)'),
    ("Best trade $", f"=MAX({T}!$G$2:$G${LAST})", '#,##0.00'),
    ("Worst trade $", f"=MIN({T}!$G$2:$G${LAST})", '#,##0.00;(#,##0.00)'),
    ("Trial target (trades)", 30, "0"),
    ("Trades remaining", f"=MAX(0,$B$12-$B$2)", "0"),
    ("Return on starting equity", f"=$B$6/{START_EQUITY:g}", "0.00%"),
]
for r, (label, val, fmt) in enumerate(items, 1):
    lc = s.cell(row=r, column=1, value=label)
    lc.font = F(bold=(r == 1))
    if val is not None:
        vc = s.cell(row=r, column=2, value=val)
        vc.font = INPUT_FONT if isinstance(val, (int, float)) else F()
        if fmt:
            vc.number_format = fmt
s["A16"] = ("Legend: blue cells are inputs — append new trades on the Trade Log sheet "
            f"(cols A-F + K); white cells are formulas. Rows pre-filled to {N_ROWS} trades.")
s["A16"].font = F(italic=True, size=9)
s["A17"] = "Source: state/trade_history.csv (Alpaca paper acct), exit types from logs/orb_*.log; trial start 2026-07-09."
s["A17"].font = F(italic=True, size=9)
s.column_dimensions["A"].width = 24
s.column_dimensions["B"].width = 14

# --- Trial plan / schedule (agreed 2026-07-24) ---
PLAN = [
    ("TRIAL PLAN", None),
    ("Jul 27 - 31", "Finish trial week on current config (v1.7) - no changes, even past trade 30"),
    ("Aug 1 - 2", "Tuning weekend: 1-2 changes max, from MFE/MAE data + sweep.py run (exit asymmetry first)"),
    ("Aug 3 - 7", "Validation week on tuned config, allow_shorts: false (mirror live cash account)"),
    ("Aug 10 ->", "If validation green/flat-with-better-exits: fund fraction of ~EUR 1,900 (cash acct)"),
]
PLAN_ROW = 19
for i, (when, what) in enumerate(PLAN):
    r = PLAN_ROW + i
    wc = s.cell(row=r, column=1, value=when)
    wc.font = F(bold=True) if i == 0 else F(bold=True, size=9)
    if what:
        dc = s.cell(row=r, column=2, value=what)
        dc.font = F(size=9)

chart = LineChart()
chart.title = "Cumulative P&L ($)"
chart.height, chart.width = 8, 16
data = Reference(ws, min_col=10, min_row=1, max_row=len(rows) + 1)
cats = Reference(ws, min_col=2, min_row=2, max_row=len(rows) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
s.add_chart(chart, "D2")

wb.save(OUT)
print(f"trade_log.xlsx regenerated: {len(rows)} trades")
